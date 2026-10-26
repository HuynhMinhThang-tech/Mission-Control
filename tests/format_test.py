# -*- coding: utf-8 -*-
"""Kiểm thử hồi quy cho: số/đơn vị tiếng Việt, nhãn cột, tiêu đề biểu đồ khớp dữ liệu, nhiều sheet, nhà cung cấp LLM.
Dựng lại ĐÚNG các lỗi người dùng đã thấy (2.24 trillion VNĐ, 11773398767.5, 'điểm', trongLuongKg, tiêu đề 'giảm 15,7%' trên biểu đồ tăng...).
Chạy: python tests/format_test.py"""
import io
import os
import re
import sys
import tempfile
import zipfile

tmp = tempfile.mkdtemp()
os.environ.update(DATABASE_URL=f"sqlite:///{tmp}/f.db")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from app import create_app, db, labels, numfmt, providers, state  # noqa: E402
from app.data.loader import list_sheets, load_files  # noqa: E402
from app.export import publish  # noqa: E402

n_ok = 0


def check(cond, label):
    global n_ok
    if not cond:
        print("FAIL:", label)
        sys.exit(1)
    n_ok += 1
    print("ok  ", label)


# ------------------------------------------------------------------ 1. số và đơn vị --
known = numfmt.Known.from_dfs([pd.DataFrame({"Tháng": ["9", "8"], "Doanh thu": [2240813537.5, 1935760431.25]})])
t = numfmt.fix_text("Doanh thu tháng 9 đạt 2.24 trillion VNĐ, cao nhất trong năm.", known)
check("trillion" not in t and "VNĐ" not in t and "2,24 tỷ đồng" in t, f"'2.24 trillion VNĐ' -> '2,24 tỷ đồng' ({t})")
t = numfmt.fix_text("Doanh thu tháng 9 lên tới 2 241 813 537 điểm, vượt 12 % so với tháng 8.", known)
check("2.241.813.537 đồng" in t and "12%" in t and "điểm" not in t, f"số ngăn bằng dấu cách + 'điểm' -> chuẩn Việt ({t})")
t = numfmt.fix_text("Tháng 9 = 2240813537.5, tỷ lệ tăng 2240813537.5/1935760431.25-1≈0.157", known)
check("2.240.813.537,5" in t and "0,157" in t and "." not in t.replace("2.240.813.537", "").replace("1.935.760.431", ""), f"số thô có dấu ngăn nghìn + dấu phẩy thập phân ({t})")
t = numfmt.fix_text("Tỷ lệ hoàn trả cao nhất tháng 3 (6.97%), vượt mức 4.9%. Ngày 24.10.2026, bản 1.2.3, giá 2.240 đồng.")
check("6,97%" in t and "4,9%" in t and "24.10.2026" in t and "1.2.3" in t and "2.240 đồng" in t, f"đổi dấu thập phân nhưng không đụng ngày/phiên bản/số nghìn ({t})")
check(numfmt.fix_text(numfmt.fix_text("Tổng 2.24 trillion VNĐ, 6.97%", known), known) == numfmt.fix_text("Tổng 2.24 trillion VNĐ, 6.97%", known), "chuẩn hóa là idempotent (chạy hai lần không đổi)")
kp = numfmt.fix_kpi({"label": "Doanh thu tổng năm 2026", "value": "11773398767.5"})
check(kp["value"] == "11,77 tỷ đồng", f"KPI 11773398767.5 -> {kp['value']}")
check(numfmt.fix_kpi({"label": "Số đơn", "value": "14500"})["value"] == "14.500", "KPI số đếm có dấu ngăn nghìn")
check(numfmt.fix_kpi({"label": "Tỷ lệ hoàn trả", "value": "0.0697"})["value"] == "6,97%", "KPI tỷ lệ 0-1 -> phần trăm")
check(not numfmt.has_value(numfmt.fix_kpi({"label": "Tỷ trọng bán hàng theo vùng", "value": ""})), "KPI rỗng bị nhận diện để loại")
bad = numfmt.audit("Điện thoại & phụ kiện chiếm 1.557.676.062 đồng, tháng 9 đạt 2,24 tỷ đồng", numfmt.Known([2557676062.5, 2240813537.5]))
check(bad == ["1.557.676.062"], f"bắt số lớn không khớp bảng ({bad}), không bắt số đúng")

check(numfmt.audit("Doanh thu tháng 9 lên tới 2.241.813.537 đồng", numfmt.Known([2240813537.5, 1935760431.25])) == ["2.241.813.537"],
      "bắt số AI viết lệch 1 triệu so với bảng (2.241.813.537 so với 2.240.813.537,5)")
check(numfmt.audit("Tăng 304.053.106 đồng so với tháng 8", numfmt.Known.from_dfs([pd.DataFrame({"Doanh thu": [2240813537.5, 1936760431.25]})])) == [],
      "số suy ra hợp lệ (chênh lệch hai ô cùng cột) không bị báo oan")

# ------------------------------------------------------------------ 1b. tín hiệu phạm vi & chất lượng (generic) --
from app.data.scope import scope_hints  # noqa: E402
_r = np.random.default_rng(2)
_n = 200
_df = pd.DataFrame({"NgayDat": pd.date_range("2026-01-01", "2026-09-24", periods=_n), "TrangThaiDon": _r.choice(["Hoàn thành", "Hoàn trả", "Đã hủy"], _n, p=[.82, .07, .11]),
                    "SoTien": np.where(_r.random(_n) < .07, -1, 1) * _r.integers(100000, 900000, _n), "ThuongHieu": _r.choice(list("ABCD"), _n), "ChiTiet": [None] * 160 + ["x"] * 40})
_h = scope_hints({"don_hang": _df})
check("«TrangThaiDon»" in _h and "QUYẾT ĐỊNH" in _h, "phát hiện cột phân loại bản ghi cần quyết định phạm vi")
check("«SoTien»" in _h and "âm" in _h, "phát hiện cột số vừa âm vừa dương")
check("24/30" in _h and "tạm tính" in _h, "phát hiện kỳ cuối chưa đủ ngày")
check("«ChiTiet» thiếu 80%" in _h, "phát hiện cột thiếu nhiều không dùng để chia cơ cấu")
check("«ThuongHieu»" not in _h, "cột phân loại thường (thương hiệu) không bị coi là loại bản ghi")

# ------------------------------------------------------------------ 2. nhãn và đơn vị cột --
for raw, want in [("DanhMuc", "Danh mục"), ("ThuongHieu", "Thương hiệu"), ("TrongLuongKg", "Trọng lượng"), ("GiaNiemYet", "Giá niêm yết"),
                  ("GiaVon", "Giá vốn"), ("order_id", "Mã đơn hàng"), ("Doanh thu (đ)", "Doanh thu (đ)"), ("PhuongThucThanhToan", "Phương thức thanh toán"), ("Customer_Name", "Tên khách hàng"),
                  ("du_lieu_thuong_mai_dien_tu", "Dữ liệu thương mại điện tử")]:
    check(labels.label(raw) == want, f"nhãn: {raw} -> {labels.label(raw)}")
check(labels.unit("TrongLuongKg") == "kg" and labels.unit("ThuongHieu") == "" and labels.unit("SoLuong") == "" and labels.unit("GiaNiemYet") == "đ"
      and labels.unit("DoanhThu") == "đ" and labels.unit("TyLeHoanTra") == "%", "đơn vị suy theo TỪ: kg không phải tiền, thương hiệu không phải thưởng")
check(labels.is_rate("GiaNiemYet") and labels.is_rate("TrongLuongKg") and not labels.is_rate("SoLuong") and labels.is_id("MaSanPham") and not labels.is_id("DanhMuc"), "nhận diện đơn giá/tỷ lệ (dùng trung bình) và cột mã")

# ------------------------------------------------------------------ 3. tiêu đề biểu đồ khớp dữ liệu --
rows = [("2026-01-01", 1.2e9), ("2026-02-01", 1.0e9), ("2026-09-01", 2.2e9)]
check(not numfmt.consistent("Doanh thu giảm 15,7% từ tháng 1 đến 3", rows, True), "tiêu đề 'giảm' trên biểu đồ tăng bị phát hiện")
check(numfmt.consistent("Doanh thu tăng 83,3% từ đầu kỳ", rows, True), "tiêu đề đúng hướng, đúng % được giữ")
check(numfmt.headline(rows, "line", "Doanh thu (đ)", True) == "Doanh thu tăng 83,3% từ T1/2026 đến T9/2026", f"câu tính từ dữ liệu: {numfmt.headline(rows, 'line', 'Doanh thu (đ)', True)}")

# ------------------------------------------------------------------ 4. toàn bộ báo cáo: dashboard, HTML, MD, PPTX --
app = create_app()
u = db.create_user("An", "an@gmail.com", "matkhau-123", status="approved")
rng = np.random.default_rng(7)
n = 69
prod = pd.DataFrame({"MaSanPham": [f"SP{i:03d}" for i in range(n)], "DanhMuc": rng.choice(["Điện thoại & Phụ kiện", "Thời trang", "Gia dụng"], n),
                     "ThuongHieu": rng.choice(["Alpha", "Beta", "Gamma", "Delta"], n), "TrongLuongKg": rng.uniform(.5, 5, n).round(1),
                     "GiaNiemYet": rng.choice([1_200_000, 2_500_000, 3_900_000], n), "GiaVon": rng.choice([700_000, 1_500_000, 2_800_000], n)})
days = pd.date_range("2026-01-01", "2026-09-24", freq="D")
sales = pd.DataFrame({"NgayDat": days, "DoanhThu": (rng.uniform(3e7, 9e7, len(days)) * np.linspace(1, 2.2, len(days))).round(0)})
month = pd.DataFrame({"Tháng": [f"2026-{m:02d}-01" for m in range(1, 10)], "Doanh thu": [1.2e9, 1.0e9, 1.1e9, 1.05e9, 1.15e9, 1.25e9, 1.5e9, 1.9357604e9, 2.2408135e9]})
region = pd.DataFrame({"Vùng": ["Miền Nam", "Miền Bắc", "Miền Trung"], "Doanh thu": [5759948862.0, 4689363481.25, 2211528975.0]})


def tab(title, df):
    return {"title": title, "cols": list(df.columns), "rows": df.values.tolist(), "err": None}


with app.test_request_context():
    state.bind(u["id"])
    A = state.A
    A["question"], A["ctx"], A["started"] = "Cho tôi tình hình kinh doanh, sản phẩm bán chạy", "Thương mại điện tử", state.now_iso()
    A["files"] = [{"name": "du_lieu_thuong_mai_dien_tu.xlsx::SanPham", "rows": n}, {"name": "du_lieu_thuong_mai_dien_tu.xlsx::DoanhThu", "rows": len(days)}]
    A["clean"] = {"du_lieu_thuong_mai_dien_tu.xlsx::SanPham": prod, "du_lieu_thuong_mai_dien_tu.xlsx::DoanhThu": sales}
    A["x"] = {3: {0: month, 1: region}}
    A["d"] = {1: {"kpis": []}, 2: {"pv": None},
              3: {"phase": "result", "tasks": [{"title": "Doanh thu hàng tháng", "goal": "xu hướng"}, {"title": "Doanh thu theo vùng", "goal": "so sánh"}],
                  "result": {"kpis": [{"label": "Doanh thu tổng năm 2026", "value": "11773398767.5"}, {"label": "Doanh thu bán chạy nhất", "value": "2258413537.5"},
                                      {"label": "Tỷ trọng bán hàng theo vùng", "value": ""}],
                             "tables": [tab("Doanh thu hàng tháng", month), tab("Doanh thu theo vùng", region)], "findings": []}},
              4: {"summary": "Doanh thu tháng 9 đạt 2.24 trillion VNĐ, cao nhất trong năm, với điện thoại & phụ kiện chiếm 57% tổng doanh thu.",
                  "insights": [{"title": "Doanh thu tháng 9 lên tới 2 241 813 537 điểm, vượt 12 % so với tháng 8", "confidence": "Cao", "source": "Doanh thu hàng tháng",
                                "evidence": "Tháng 9 = 2240813537.5, Tháng 8 = 1935760431.25, tỷ lệ tăng 2240813537.5/1935760431.25-1≈0.157"},
                               {"title": "Điện thoại & phụ kiện chiếm 1 557 676 062 điểm, tương đương 57% doanh thu", "confidence": "Cao", "source": "Doanh thu theo vùng",
                                "evidence": "Điện thoại & phụ kiện = 2557676062.5"}],
                  "actions": [{"text": "Tối ưu khuyến mãi", "priority": "Cao", "impact": "Tăng 1.5 triệu VND"}], "risks": ["Tháng 9 mới có 24 ngày"]},
              5: {"title": "Tình hình doanh thu 2026", "narrative": "Doanh thu tháng 9 đạt 2.24 trillion VNĐ.",
                  "final": {"title": "Tình hình doanh thu 2026", "kpis": [0, 1, 2],
                            "charts": [{"task": 0, "x": "Tháng", "y": "Doanh thu", "kind": "line", "title": "Doanh thu tháng 2026", "unit": "đ",
                                        "insight": "Doanh thu giảm 15,7% từ tháng 1 đến 3", "recommended": True},
                                       {"task": 1, "x": "Vùng", "y": "Doanh thu", "kind": "bars", "title": "Doanh thu theo vùng", "unit": "đ", "insight": "", "recommended": True}]}}}
    publish()
    aid = A["id"]
    state.unbind()

get = lambda name: db.get_file(u["id"], aid, name)[0]       # noqa: E731
html, md, dash = get("bao_cao.html").decode("utf8"), get("bao_cao.md").decode("utf8"), get("dashboard.html").decode("utf8")
from pptx import Presentation  # noqa: E402
prs = Presentation(io.BytesIO(get("bao_cao.pptx")))
ppt = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
for name, body in [("báo cáo HTML", html), ("Markdown", md), ("PPTX", ppt), ("dashboard", dash)]:
    check(not re.search(r"trillion|billion|million|VNĐ|VND", body.replace("VND_", "")) , f"{name}: không còn trillion/VND tiếng Anh")
    check("11773398767.5" not in body and "2240813537.5" not in body and "1935760431.25" not in body, f"{name}: không còn số thô thiếu dấu ngăn nghìn")
check("11,77 tỷ đồng" in html and "11,77 tỷ đồng" in md, "KPI hiện '11,77 tỷ đồng' ở HTML và Markdown")
check("Tỷ trọng bán hàng theo vùng" not in html and "Tỷ trọng bán hàng theo vùng" not in ppt, "KPI rỗng không lên báo cáo và PPTX")
check("giảm 15,7%" not in html and "giảm 15,7%" not in ppt and "tăng" in html and "T1/2026" in html, "tiêu đề biểu đồ sai hướng được thay bằng câu tính từ dữ liệu")
check("điểm" not in html.replace("Điểm", "") .replace("điểm bất", ""), "đơn vị tiền không còn là 'điểm'")
check("Cần kiểm tra số liệu" in html, "phát hiện có số lớn không khớp bảng được gắn cờ 'Cần kiểm tra số liệu'")
check("Danh mục" in dash and "Thương hiệu" in dash and "Trọng lượng" in dash and '"DanhMuc"' not in dash and "trongLuongKg" not in dash, "dashboard dùng nhãn tiếng Việt thay tên biến")
check("class=\"dtabs\"" in dash and "data-ds=" in dash and 'id="dsw"' not in dash, "dashboard đổi bảng bằng tab, không còn ô chọn")

# ------------------------------------------------------------------ 5. nhiều sheet = nhiều tab --
buf = io.BytesIO()
with pd.ExcelWriter(buf, engine="openpyxl") as w:
    pd.DataFrame({"Thang": ["2026-01", "2026-02", "2026-03"], "DoanhThu": [100, 120, 90]}).to_excel(w, sheet_name="DoanhThu", index=False)
    pd.DataFrame({"DanhMuc": list("ABCD"), "SoLuong": [3, 4, 5, 6]}).to_excel(w, sheet_name="SanPham", index=False)


class F:
    filename = "ban_hang.xlsx"

    def read(self):
        return buf.getvalue()


info = list_sheets([F()])[0]
check(info["sheets"] == ["DoanhThu", "SanPham"] and info["info"]["SanPham"]["rows"] == 5, "liệt kê sheet kèm số dòng x cột")
d, *_ = load_files([F()], {"ban_hang.xlsx": ["DoanhThu", "SanPham"]})
check(list(d) == ["ban_hang.xlsx::DoanhThu", "ban_hang.xlsx::SanPham"], "chọn nhiều sheet -> mỗi sheet một bảng 'tệp::sheet'")
d, *_ = load_files([F()], {"ban_hang.xlsx": "SanPham"})
check(list(d) == ["ban_hang.xlsx::SanPham"], "vẫn nhận chọn một sheet dạng chuỗi (tương thích cũ)")

# ------------------------------------------------------------------ 6. nhà cung cấp và Base URL --
check(providers.check_pair("https://api.openai.com/v1", "gemini-2.5-flash") and "Google Gemini" in providers.check_pair("https://api.openai.com/v1", "gemini-2.5-flash"),
      "model Gemini + Base URL OpenAI: báo lỗi tiếng Việt kèm cách sửa")
check(providers.check_pair("https://generativelanguage.googleapis.com/v1beta/openai/", "gemini-2.5-flash") is None and providers.check_pair("https://openrouter.ai/api/v1", "gpt-4o") is None
      and providers.check_pair("https://api.noi-bo.vn/v1", "gpt-4o") is None, "khớp hoặc là gateway/địa chỉ riêng thì không chặn")
check(providers.BY_ID["gemini"]["base_url"].startswith("https://generativelanguage.googleapis.com/v1beta/openai") and providers.provider_of_url("https://api.anthropic.com/v1/") == "anthropic",
      "mỗi nhà cung cấp có Base URL mặc định riêng; nhận ra nhà cung cấp từ địa chỉ")

print(f"\nTẤT CẢ {n_ok} KIỂM TRA ĐỀU ĐẠT")

# -*- coding: utf-8 -*-
"""Bản trình chiếu PPTX theo phong cách consultant (16:9, answer-first):
bìa → tóm tắt điều hành → chỉ số → mỗi biểu đồ 1 slide với tiêu đề là KẾT LUẬN + "so what"
→ insight → khuyến nghị → rủi ro & quyết định cần có → phụ lục (phương pháp, dữ liệu)."""
import io

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE as CT, XL_LABEL_POSITION as LP, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..utils import compact, vn
from .common import PRI_ORDER, chart_rows, facts, scale, trunc

NAVY, PRI, SOFT, INK, SUB, LINE, BG = "0F172A", "4F46E5", "EEF2FF", "1F2937", "5B6275", "E5E7EB", "F8FAFC"
OK, WARN, BAD, MUTE, WHITE = "0F8F6A", "B76A00", "D6304A", "B4BAD6", "FFFFFF"
FONT = "Arial"   # một họ chữ duy nhất cho toàn bộ deck (hỗ trợ tiếng Việt tốt, có sẵn trên Windows/Mac)
# Thang cỡ chữ thống nhất
SZ_TITLE, SZ_BODY, SZ_SMALL, SZ_CAP = 24, 14, 12, 10
DONUT_PAL = ["4F46E5", "0EA5E9", "10B981", "F59E0B", "EF4444", "8B5CF6", "64748B"]
W, H, MX = 13.333, 7.5, 0.6
CW = W - 2 * MX


# ----------------------------------------------------------------- helpers --
def rgb(h):
    return RGBColor.from_string(h)


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(.75)
    else:
        sh.line.fill.background()
    return sh


def text(s, x, y, w, h, content, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=0):
    """content: str | list[str | (str, dict)]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(.05)
    tf.margin_top = tf.margin_bottom = Inches(.03)
    for i, p in enumerate(content if isinstance(content, list) else [content]):
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = o.get("align", align)
        para.space_after = Pt(o.get("after", after))
        r = para.add_run()
        r.text = t
        f = r.font
        f.name, f.size = FONT, Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = rgb(o.get("color", color))
    return tb


def _ln(tcPr, side, color=None):
    el = etree.SubElement(tcPr, qn(f"a:ln{side}"))
    if color:
        el.set("w", "9525")
        sf = etree.SubElement(el, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
    else:
        el.set("w", "0")
        etree.SubElement(el, qn("a:noFill"))


def _cell(cell, txt, size=12, bold=False, color=INK, fill=WHITE, align=PP_ALIGN.LEFT):
    tcPr = cell._tc.get_or_add_tcPr()
    for sd in ("L", "R", "T"):
        _ln(tcPr, sd)
    _ln(tcPr, "B", LINE)            # viền phải được thêm TRƯỚC khi đặt màu nền (đúng thứ tự schema)
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.margin_left = cell.margin_right = Inches(.1)
    cell.margin_top = cell.margin_bottom = Inches(.05)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = txt
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
    r.font.color.rgb = rgb(color)


def table(s, x, y, colw, header, rows, size=12, rowh=.62, styles=None):
    """styles: callable(row_idx, col_idx, value) -> dict(color, fill, bold, align) | None"""
    shp = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(sum(colw)), Inches(rowh * (len(rows) + 1)))
    t = shp.table
    t.horz_banding = False
    t.first_row = True
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(cw)
    for j, h in enumerate(header):
        _cell(t.cell(0, j), h, size=size - 1, bold=True, color=SUB, fill=BG)
    t.rows[0].height = Inches(.42)
    for i, r in enumerate(rows):
        t.rows[i + 1].height = Inches(rowh)
        for j, v in enumerate(r):
            st = (styles(i, j, v) if styles else None) or {}
            _cell(t.cell(i + 1, j), str(v), size=size, **st)
    return t


class Deck:
    def __init__(self, F):
        self.F = F
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = Inches(W), Inches(H)
        self.n = 0

    def slide(self, title, kicker=None):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.n += 1
        rect(s, MX, .55, .08, .9, fill=PRI)
        if kicker:
            text(s, MX + .2, .38, CW - .2, .3, kicker.upper(), size=11, bold=True, color=PRI)
        text(s, MX + .2, .66, CW - .2, .95, trunc(title, 125), size=SZ_TITLE, bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP)
        rect(s, MX, 7.0, CW, .01, fill=LINE)
        text(s, MX, 7.05, CW - 1, .3, f"{trunc(self.F['title'], 70)}   |   {self.F['date']}", size=SZ_CAP, color=SUB)
        text(s, W - MX - 1, 7.05, 1, .3, str(self.n), size=10, color=SUB, align=PP_ALIGN.RIGHT)
        return s

    def source(self, s, msg):
        text(s, MX, 6.62, CW, .3, "Nguồn: " + msg, size=10, color=SUB)


# ------------------------------------------------------------------ slides ---
def cover(d):
    F = d.F
    s = d.p.slides.add_slide(d.p.slide_layouts[6])
    d.n += 1
    rect(s, 0, 0, W, H, fill=NAVY)
    rect(s, MX, 2.0, .12, 2.0, fill=PRI)
    text(s, MX + .35, 1.9, 10.5, 2.3, trunc(F["title"], 120), size=38, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + .35, 4.45, 10.5, 1.0, "Câu hỏi: " + trunc(F["question"], 220), size=17, color="C7CCF5")
    meta = [F["date"]] + ([F["period"]] if F.get("period") else [])
    text(s, MX + .35, 6.3, 11, .4, "   |   ".join(meta) + "   |   Dựa trên " + ", ".join(trunc(f["name"], 28) for f in F["files"][:3]), size=12, color="9CA3D9")


def exec_summary(d):
    F = d.F
    first = F["summary"].replace("\n", " ").split(". ")[0].strip().rstrip(".")
    s = d.slide(trunc(first, 130), "Tóm tắt điều hành")
    lw = 7.6
    rect(s, MX, 1.85, lw, 1.95, fill=SOFT)
    text(s, MX + .2, 1.95, lw - .4, 1.8, trunc(F["summary"], 420), size=15, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX, 4.0, lw, .35, "Những phát hiện chính", size=13, bold=True, color=PRI)
    items = [f"{i + 1}.  {trunc(x['title'], 135)}" for i, x in enumerate(F["insights"][:3])]
    text(s, MX, 4.35, lw, 2.2, items, size=13, color=INK, after=7)
    kx = MX + lw + .4
    kw = W - MX - kx
    text(s, kx, 1.85, kw, .35, "Con số cần nhớ", size=13, bold=True, color=PRI)
    ks = F["kpis"][:4]
    ch = 4.6 / max(len(ks), 1)
    for i, k in enumerate(ks):
        y = 2.25 + i * min(ch, 1.15)
        rect(s, kx, y, kw, min(ch, 1.15) - .12, fill=WHITE, line=LINE)
        text(s, kx + .15, y + .04, kw - .3, .55, trunc(k["value"], 22), size=22, bold=True, color=PRI)
        text(s, kx + .15, y + .56, kw - .3, .4, trunc(k["label"] + (" · " + k["delta"] if k.get("delta") else ""), 46), size=11, color=SUB)
    d.source(s, "Mission Control, dữ liệu đã làm sạch")


def kpi_slide(d):
    F = d.F
    ks = F["kpis"][:8]
    if not ks:
        return
    s = d.slide("Các chỉ số then chốt của kỳ phân tích", "Bức tranh số liệu")
    cols = 4 if len(ks) > 3 else len(ks)
    gap = .25
    cw = (CW - gap * (cols - 1)) / cols
    rows = (len(ks) + cols - 1) // cols
    chh = min(2.2, (4.7 - gap * (rows - 1)) / rows)
    for i, k in enumerate(ks):
        x, y = MX + (i % cols) * (cw + gap), 1.95 + (i // cols) * (chh + gap)
        rect(s, x, y, cw, chh, fill=WHITE, line=LINE)
        rect(s, x, y, cw, .07, fill=PRI)
        text(s, x + .15, y + .25, cw - .3, .8, trunc(k["value"], 20), size=28, bold=True, color=PRI, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + .15, y + 1.05, cw - .3, .5, trunc(k["label"], 40), size=13, bold=True, color=INK)
        if k.get("delta"):
            dc = OK if str(k["delta"]).strip().startswith("+") else BAD if str(k["delta"]).strip()[:1] in "-−" else SUB
            text(s, x + .15, y + 1.5, cw - .3, .4, trunc(k["delta"], 40), size=12, color=dc)
    d.source(s, "Phân tích trên dữ liệu đã làm sạch")


def _chart(s, x, y, w, h, rows, kind, div):
    cd = CategoryChartData()
    cd.categories = [trunc(l, 18) for l, _ in rows]
    cd.add_series("Giá trị", [round(v / div, 3) for _, v in rows])
    donut = kind == "donut"
    line = kind == "line"
    ctype = CT.DOUGHNUT if donut else CT.LINE_MARKERS if line else CT.COLUMN_CLUSTERED if kind == "cols" else CT.BAR_CLUSTERED
    ch = s.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_title = False
    ch.font.size, ch.font.name = Pt(SZ_SMALL), FONT
    ch.font.color.rgb = rgb(INK)
    pl = ch.plots[0]
    ser = pl.series[0]
    if donut:
        from pptx.enum.chart import XL_LEGEND_POSITION
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        ch.legend.font.size, ch.legend.font.name = Pt(SZ_SMALL), FONT
        for i in range(len(rows)):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(DONUT_PAL[i % len(DONUT_PAL)])
        pl.has_data_labels = True
        dl = pl.data_labels
        dl.show_percentage, dl.show_value, dl.show_category_name = True, False, False
        dl.number_format, dl.number_format_is_linked = "0%", False
        dl.font.size, dl.font.bold, dl.font.name = Pt(SZ_SMALL), True, FONT
        dl.font.color.rgb = rgb(WHITE)
        return
    ch.has_legend = False
    mx = max(abs(v) / div for _, v in rows) or 1
    fmt = "#,##0" if mx >= 100 else "#,##0.0" if mx >= 10 else "#,##0.00"
    hi = max(range(len(rows)), key=lambda i: abs(rows[i][1]))
    if line:
        ser.smooth = False
        ser.format.line.color.rgb = rgb(PRI)
        ser.format.line.width = Pt(2.5)
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 7
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb(PRI)
        if len(rows) <= 12:
            pl.has_data_labels = True
            pl.data_labels.position = LP.ABOVE
    else:
        pl.gap_width = 55
        pl.vary_by_categories = False
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(MUTE)
        for i, (_, v) in enumerate(rows):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(BAD if v < 0 else PRI if i == hi else MUTE)
        pl.has_data_labels = True
        pl.data_labels.position = LP.OUTSIDE_END
    if pl.has_data_labels:
        dl = pl.data_labels
        dl.font.size, dl.font.bold, dl.font.name = Pt(SZ_SMALL), True, FONT
        dl.number_format, dl.number_format_is_linked = fmt, False
    va, ca = ch.value_axis, ch.category_axis
    va.visible = False
    va.has_major_gridlines = False
    ca.has_major_gridlines = False
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(SZ_SMALL), FONT
    ca.format.line.color.rgb = rgb(LINE)
    if kind == "bars":
        ca.reverse_order = True


def chart_slides(d):
    F = d.F
    for c in F["charts"]:
        rows = chart_rows(F["dfs"][c["task"]], c["x"], c["y"], 12 if c["kind"] in ("bars", "donut") else 24)
        if not rows:
            continue
        s = d.slide(c.get("insight") or c["title"], "Phân tích")
        div, ulab = scale(rows) if c["kind"] != "donut" else (1.0, "")
        unit = " ".join(x for x in [ulab, c.get("unit", "")] if x)
        text(s, MX, 1.7, 8.0, .35, c["title"] + (f"  (đơn vị: {unit})" if unit else ""), size=12, bold=True, color=SUB)
        _chart(s, MX, 2.05, 8.0, 4.5, rows, c["kind"], div)
        px, pw = 8.95, W - MX - 8.95
        rect(s, px, 1.75, pw, 4.8, fill=SOFT)
        text(s, px + .2, 1.9, pw - .4, .35, "ĐIỀU NÀY CÓ NGHĨA LÀ", size=11, bold=True, color=PRI)
        text(s, px + .2, 2.3, pw - .4, 1.7, trunc(c.get("insight") or c["title"], 230), size=14, bold=True, color=NAVY)
        text(s, px + .2, 4.1, pw - .4, .3, "Số liệu nổi bật", size=11, bold=True, color=PRI)
        text(s, px + .2, 4.4, pw - .4, 2.1, [f"▸ {x}" for x in facts(rows, c["kind"])], size=12, color=INK, after=6)
        d.source(s, f"{trunc(F['table_titles'][c['task']], 80)}; dữ liệu đã làm sạch")


def insight_slides(d):
    ins = d.F["insights"]
    for k in range(0, len(ins), 3):
        part = ins[k:k + 3]
        s = d.slide(f"{len(ins)} phát hiện chính, mỗi phát hiện đều có bằng chứng truy ngược về dữ liệu" if k == 0 else "Các phát hiện chính (tiếp)", "Insight")
        rows = [[f"{k + i + 1}", trunc(x["title"], 190), trunc(x.get("evidence", ""), 260), x.get("confidence", "")] for i, x in enumerate(part)]
        conf = {"Cao": OK, "Trung bình": WARN, "Thấp": BAD}

        def st(i, j, v):
            if j == 0:
                return dict(bold=True, color=PRI, align=PP_ALIGN.CENTER)
            if j == 1:
                return dict(bold=True, color=NAVY)
            if j == 3:
                return dict(bold=True, color=conf.get(v, SUB), align=PP_ALIGN.CENTER)
        table(s, MX, 1.85, [.5, 5.0, 5.4, 1.23], ["#", "Insight", "Bằng chứng", "Độ tin cậy"], rows, size=SZ_SMALL, rowh=1.3, styles=st)
        d.source(s, "bảng kết quả phân tích; chi tiết ở phụ lục")


def action_slides(d):
    F = d.F
    acts = sorted(F["actions"], key=lambda a: PRI_ORDER.get(a.get("priority"), 9))
    for k in range(0, len(acts), 5):
        part = acts[k:k + 5]
        top = trunc(acts[0]["text"], 95) if k == 0 else "Khuyến nghị (tiếp)"
        s = d.slide(("Ưu tiên hàng đầu: " + top) if k == 0 else top, "Khuyến nghị")
        rows = [[a.get("priority", ""), trunc(a["text"], 170), trunc(a.get("impact", ""), 170)] for a in part]
        pc = {"Cao": (BAD, "FDE7EA"), "Trung bình": (WARN, "FCF0D9"), "Thấp": (SUB, "F1F2F6")}

        def st(i, j, v):
            if j == 0:
                c, f = pc.get(v, (SUB, WHITE))
                return dict(bold=True, color=c, fill=f, align=PP_ALIGN.CENTER)
            if j == 1:
                return dict(bold=True, color=NAVY)
        table(s, MX, 1.85, [1.4, 5.8, 4.93], ["Ưu tiên", "Hành động", "Tác động kỳ vọng"], rows, size=SZ_BODY - 1, rowh=.9, styles=st)
        d.source(s, "insight đã được xác nhận ở bước trước")


def decision_slide(d):
    F = d.F
    s = d.slide("Cần lãnh đạo quyết định các hạng mục ưu tiên cao và theo dõi các rủi ro sau", "Quyết định và rủi ro")
    hw = (CW - .4) / 2
    rect(s, MX, 1.85, hw, 3.9, fill=SOFT)
    text(s, MX + .2, 1.95, hw - .4, .35, "ĐỀ NGHỊ PHÊ DUYỆT", size=11, bold=True, color=PRI)
    top = [a for a in F["actions"] if a.get("priority") == "Cao"][:4] or F["actions"][:3]
    text(s, MX + .2, 2.35, hw - .4, 4.0, [f"{i + 1}.  {trunc(a['text'], 140)}" for i, a in enumerate(top)] or ["—"], size=14, color=NAVY, after=10)
    x2 = MX + hw + .4
    rect(s, x2, 1.85, hw, 3.9, fill=WHITE, line=LINE)
    text(s, x2 + .2, 1.95, hw - .4, .35, "RỦI RO CẦN THEO DÕI", size=11, bold=True, color=BAD)
    text(s, x2 + .2, 2.35, hw - .4, 4.0, [f"▸ {trunc(r, 150)}" for r in F["risks"][:5]] or ["Chưa ghi nhận rủi ro đáng kể."], size=13, color=INK, after=8)
    d.source(s, "tổng hợp từ insight và hành động")


def appendix(d):
    F = d.F
    s = d.slide("Phụ lục: dữ liệu, cách làm sạch và các phân tích đã chạy", "Phương pháp")
    colw = (CW - .6) / 3
    q = F.get("quality") or {}
    qb = sum((q.get("before") or {}).values())
    qa = sum((q.get("after") or {}).values())
    blocks = [
        ("DỮ LIỆU", [f"{trunc(f['name'], 34)}: {vn(f['rows'], 0)} dòng" for f in F["files"][:5]] + ([f"Khoảng thời gian: {F['period']}"] if F.get("period") else [])),
        ("LÀM SẠCH", ([f"Vấn đề trước/sau: {vn(qb, 0)} → {vn(qa, 0)}"] if q else []) + [trunc(x, 60) for x in F["applied"][:6]] or ["Không áp dụng thao tác nào"]),
        ("PHÂN TÍCH", [trunc(t["title"], 70) for t in F["tasks"][:7]]),
    ]
    for i, (h, lines) in enumerate(blocks):
        x = MX + i * (colw + .3)
        rect(s, x, 1.85, colw, 3.9, fill=BG, line=LINE)
        text(s, x + .15, 1.95, colw - .3, .35, h, size=11, bold=True, color=PRI)
        text(s, x + .15, 2.35, colw - .3, 4.0, [f"▸ {x_}" for x_ in lines], size=12, color=INK, after=6)
    d.source(s, "Mission Control; bằng chứng đầy đủ trong dashboard.html")


def build_pptx(F):
    d = Deck(F)
    cover(d)
    exec_summary(d)
    kpi_slide(d)
    chart_slides(d)
    if F["insights"]:
        insight_slides(d)
    if F["actions"]:
        action_slides(d)
    decision_slide(d)
    appendix(d)
    b = io.BytesIO()
    d.p.save(b)
    return b.getvalue()
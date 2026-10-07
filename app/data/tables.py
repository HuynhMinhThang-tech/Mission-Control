# -*- coding: utf-8 -*-
"""Tự nhận diện BẢNG nằm trong một sheet (kiểu Power BI / Excel Table), dù sheet không bắt đầu ở A1.

Thứ tự ưu tiên (xác định, không dùng AI):
  1. Bảng Excel thật (Insert > Table / ListObject): tên và vùng do người dùng khai báo -> tin cậy 100.
  2. Bảng suy ra từ bố cục: cắt đệ quy theo dòng trống / cột trống (XY-cut) -> các vùng chữ nhật,
     nhận dòng tiêu đề nếu có, lấy tên từ dòng tiêu đề nằm ngay phía trên, chấm điểm rồi xếp hạng.
Vùng nhỏ (ghi chú, tiêu đề trang, 1 ô) không phải bảng và bị bỏ qua.
"""
import re

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter as col_letter
from openpyxl.utils import range_boundaries

from ..utils import to_num_if_mostly
from .structure import TOTAL

MIN_DATA_ROWS = 2     # tối thiểu số dòng dữ liệu (không tính tiêu đề)
MIN_COLS = 2
TITLE_GAP = 2         # tiêu đề bảng nằm tối đa 2 dòng phía trên bảng


# ------------------------------------------------------------------ tiện ích --
def _is_text(v):
    return isinstance(v, str) and v.strip() != ""


def _is_num(v):
    return isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, (bool, np.bool_)) and not pd.isna(v)


def _runs(flags):
    """[(đầu, cuối_không_gồm)] của các đoạn True liên tiếp."""
    out, start = [], None
    for i, f in enumerate(flags):
        if f and start is None:
            start = i
        elif not f and start is not None:
            out.append((start, i))
            start = None
    if start is not None:
        out.append((start, len(flags)))
    return out


def _header_like(vals):
    """Dòng có >=2 ô, toàn chữ, không trùng nhau -> giống dòng tiêu đề."""
    v = [x for x in vals if not pd.isna(x)]
    return len(v) >= 2 and all(_is_text(x) for x in v) and len({str(x).strip().lower() for x in v}) == len(v)


def _has_numbers(vals):
    return any(_is_num(x) or isinstance(x, (pd.Timestamp,)) for x in vals)


def ref_of(r0, r1, c0, c1, off=(0, 0)):
    """Tọa độ kiểu Excel, vd 'B3:F40' (off = vị trí gốc của lưới trong sheet, tính từ 0)."""
    return f"{col_letter(c0 + off[1] + 1)}{r0 + off[0] + 1}:{col_letter(c1 - 1 + off[1] + 1)}{r1 - 1 + off[0] + 1}"


# ------------------------------------------------------- bước 1: cắt vùng -----
def _split(mask, raw, r0, r1, c0, c1, out, rows_ok=True):
    sub = mask[r0:r1, c0:c1]
    if not sub.any():
        return
    rs, cs = np.where(sub.any(1))[0], np.where(sub.any(0))[0]
    r0, r1, c0, c1 = r0 + rs[0], r0 + rs[-1] + 1, c0 + cs[0], c0 + cs[-1] + 1
    sub = mask[r0:r1, c0:c1]
    rb = _runs(sub.any(1))
    if rows_ok and len(rb) > 1:
        # gộp lại nếu dải dưới chỉ là phần nối tiếp dữ liệu của dải trên (bảng có 1 dòng trống ở giữa)
        bands = [[r0 + a, r0 + b] for a, b in rb]
        merged = [bands[0]]
        for a, b in bands[1:]:
            pa, pb = merged[-1]
            cols_a = np.where(mask[pa:pb, c0:c1].any(0))[0]
            cols_b = np.where(mask[a:b, c0:c1].any(0))[0]
            overlap = len(set(cols_a) & set(cols_b)) / max(1, min(len(cols_a), len(cols_b)))
            first = raw.iloc[a, c0:c1].tolist()
            if (pb - pa) >= 2 and overlap >= .7 and not _header_like(first) and _has_numbers(first):
                merged[-1][1] = b
                merged[-1].append("joined")
            else:
                merged.append([a, b])
        if len(merged) > 1 or len(merged[0]) > 2:
            for band in merged:
                _split(mask, raw, band[0], band[1], c0, c1, out, rows_ok=len(band) == 2)
            return
    cb = _runs(sub.any(0))
    if len(cb) > 1:
        for a, b in cb:
            _split(mask, raw, r0, r1, c0 + a, c0 + b, out, rows_ok)
        return
    out.append((r0, r1, c0, c1))


# --------------------------------------------- bước 2: dựng bảng từ một vùng ---
def _dedupe(names):
    seen, out = {}, []
    for j, n in enumerate(names):
        n = str(n).strip() if not pd.isna(n) and str(n).strip() else f"cot_{j + 1}"
        k = seen.get(n, 0)
        seen[n] = k + 1
        out.append(n if k == 0 else f"{n}_{k + 1}")
    return out


def _coerce(df):
    for c in df.columns:
        s = df[c]
        df[c] = to_num_if_mostly(s)
    return df


def _make_df(blk, header):
    """blk: vùng chữ nhật (DataFrame). Trả về (df, ghi_chú)."""
    notes = []
    if header:
        cols = _dedupe(blk.iloc[0].tolist())
        body = blk.iloc[1:].copy()
    else:
        cols = [f"cot_{j + 1}" for j in range(blk.shape[1])]
        body = blk.copy()
        notes.append("Không có dòng tiêu đề rõ ràng, đặt tên cột tạm là cot_1, cot_2…")
    body.columns = cols
    body = body.dropna(how="all").reset_index(drop=True)
    body = body.loc[:, ~(body.isna().all() & pd.Series([c.startswith("cot_") for c in body.columns], index=body.columns))]
    # dòng tổng ở cuối bảng -> tách ra để không cộng trùng
    dropped = 0
    while len(body) > 2 and _is_text(body.iat[len(body) - 1, 0]) and TOTAL.match(body.iat[len(body) - 1, 0]):
        body = body.iloc[:-1]
        dropped += 1
    if dropped:
        notes.append(f"Tách {dropped} dòng TỔNG ở cuối bảng để không cộng trùng")
    return _coerce(body.reset_index(drop=True)), notes


def _score(blk, df, has_header, source):
    if source == "excel_table":
        return 100
    s = 100
    if not has_header:
        s -= 20
    dens = float(df.notna().to_numpy().mean()) if df.size else 0
    s -= max(0, (.9 - dens)) * 60
    if len(df) < 5:
        s -= 15
    mixed = 0
    for c in df.columns:
        v = df[c].dropna()
        if len(v) and 0 < sum(_is_num(x) for x in v) < len(v) * .8 and sum(_is_text(x) for x in v) > len(v) * .1:
            mixed += 1
    s -= 25 * mixed / max(1, df.shape[1])
    return int(max(0, min(100, round(s))))


def _title_text(raw, r0, c0, c1, taken, off):
    """Dòng ghi chú/tiêu đề 1 ô ngay phía trên bảng -> dùng làm tên bảng."""
    for k in range(1, TITLE_GAP + 1):
        r = r0 - k
        if r < 0:
            break
        vals = [v for v in raw.iloc[r, c0:c1].tolist() if _is_text(v)]
        if len(vals) == 1 and (r, c0) not in taken:
            t = re.sub(r"^[▸►▶•\-\s]+", "", vals[0]).strip()
            taken.add((r, c0))
            return t[:40]
        if vals:
            break
    return None


# --------------------------------------------------------------- điểm vào ------
def detect_tables(raw, off=(0, 0)):
    """raw: lưới ô (DataFrame, header=None). Trả về danh sách bảng đã xếp hạng:
    [{name, ref, df, score, source, notes, rank, rows, cols}]"""
    if raw is None or raw.empty:
        return []
    raw = raw.replace(r"^\s*$", np.nan, regex=True).reset_index(drop=True)
    raw.columns = range(raw.shape[1])
    mask = raw.notna().to_numpy()
    leaves = []
    _split(mask, raw, 0, raw.shape[0], 0, raw.shape[1], leaves)

    found, taken, k = [], set(), 0
    for (r0, r1, c0, c1) in sorted(leaves, key=lambda x: (x[0], x[2])):
        blk = raw.iloc[r0:r1, c0:c1]
        inner_title = None
        # tiêu đề bảng dính liền ngay trên dòng tiêu đề cột (không cách dòng trống) -> tách ra làm tên bảng
        while blk.shape[0] >= 3 and blk.iloc[0].notna().sum() == 1 and _header_like(blk.iloc[1].tolist()):
            t0 = next(v for v in blk.iloc[0].tolist() if not pd.isna(v))
            inner_title = re.sub(r"^[▸►▶•\-\s]+", "", str(t0)).strip()[:40] or inner_title
            blk, r0 = blk.iloc[1:], r0 + 1
        first = blk.iloc[0].tolist()
        below = blk.iloc[1:]
        header = _header_like(first)
        if header and len(below):
            # tiêu đề thật: bên dưới có dữ liệu khác tiêu đề (không lặp lại đúng chữ của tiêu đề)
            labels = {str(x).strip().lower() for x in first if not pd.isna(x)}
            header = not any(str(x).strip().lower() in labels for x in below.to_numpy().ravel() if _is_text(x))
        n_data = (r1 - r0) - (1 if header else 0)
        if (c1 - c0) < MIN_COLS or n_data < MIN_DATA_ROWS:
            continue
        df, notes = _make_df(blk, header)
        if len(df) < MIN_DATA_ROWS or df.shape[1] < MIN_COLS:
            continue
        k += 1
        found.append({"name": inner_title or _title_text(raw, r0, c0, c1, taken, off) or f"Bảng {k}", "ref": ref_of(r0, r1, c0, c1, off), "df": df,
                      "score": _score(blk, df, header, "layout"), "source": "layout", "notes": notes, "header": header})
    return _rank(found)


def _rank(found):
    """Ưu tiên bảng lớn và sạch: điểm = chất lượng x log(số ô dữ liệu)."""
    for t in found:
        t["rows"], t["cols"] = int(len(t["df"])), int(t["df"].shape[1])
        t["_w"] = t["score"] * np.log1p(t["rows"] * t["cols"]) + (1000 if t["source"] == "excel_table" else 0)
    found.sort(key=lambda t: -t["_w"])
    names = {}
    for i, t in enumerate(found):
        t["rank"] = i + 1
        n = names.get(t["name"], 0)
        names[t["name"]] = n + 1
        if n:
            t["name"] = f"{t['name']} ({n + 1})"
        del t["_w"]
    return found


def excel_tables(ws):
    """Bảng Excel thật (ListObject) của một worksheet openpyxl. Trả về danh sách như detect_tables."""
    found = []
    try:
        items = list(ws.tables.items())
    except Exception:
        return []
    for name, tb in items:
        ref = getattr(tb, "ref", tb)
        try:
            c0, r0, c1, r1 = range_boundaries(ref)
            rows = [list(r) for r in ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1, values_only=True)]
            tot = int(getattr(tb, "totalsRowCount", 0) or 0)
            if tot:
                rows = rows[:-tot]
            if len(rows) < 1 + MIN_DATA_ROWS:
                continue
            df = pd.DataFrame(rows[1:], columns=_dedupe(rows[0])).dropna(how="all").reset_index(drop=True)
            found.append({"name": str(name), "ref": ref, "df": _coerce(df), "score": 100, "source": "excel_table", "header": True,
                          "notes": ["Bảng Excel (Table) do người dùng định nghĩa" + (", đã bỏ dòng tổng" if tot else "")]})
        except Exception:
            continue
    return _rank(found)

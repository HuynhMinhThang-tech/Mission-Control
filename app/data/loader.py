# -*- coding: utf-8 -*-
"""Đọc tệp người dùng tải lên + bộ dữ liệu mẫu (chỉ dùng khi người dùng chủ động bấm)."""
import csv
import io
import json
import re
import traceback

import numpy as np
import pandas as pd

from .structure import analyze, find_header_rows, meta
from .tables import detect_tables, excel_tables


def _open_wb(b):
    """Mở workbook openpyxl MỘT lần cho cả file (xlsx/xlsm); lỗi thì trả None."""
    try:
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(b), data_only=True)
    except Exception:
        return None


def _csv_grid(b):
    """CSV -> lưới ô GIỮ nguyên dòng trống (pandas mặc định bỏ dòng trống nên không dò được ranh giới bảng)."""
    text = b.decode("utf-8-sig", "replace")
    try:
        dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    w = max((len(r) for r in rows), default=0)
    if not w:
        return None

    def cell(x):
        x = x.strip()
        if not x:
            return np.nan
        try:
            return float(x) if re.fullmatch(r"-?\d+(\.\d+)?", x) else x
        except ValueError:
            return x
    return pd.DataFrame([[cell(c) for c in r] + [np.nan] * (w - len(r)) for r in rows])


def _sheet_grid(wb, sheet):
    """Lưới ô của sheet theo đúng tọa độ Excel. Trả về (DataFrame, (dòng_gốc, cột_gốc))."""
    ws = wb[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return (pd.DataFrame(rows) if rows else None), (ws.min_row - 1, ws.min_column - 1)


BIG = 3_000_000   # sheet quá lớn mà đã "tidy" thì bỏ qua bước dò bảng cho nhanh


def _found_tables(src, raw, res):
    """Trả về danh sách bảng nhận diện được, hoặc [] để dùng luồng cũ."""
    kind, wb, sheet, blob = src
    tidy = res["verdict"] == "tidy"
    try:
        if kind == "xlsx" and wb is not None:
            ex = excel_tables(wb[sheet])                      # ưu tiên 1: bảng Excel thật
            if ex and (len(ex) >= 2 or not tidy):
                return ex
            grid, off = _sheet_grid(wb, sheet)
        else:
            grid, off = (_csv_grid(blob) if kind == "csv" else raw), (0, 0)
        if grid is None or (tidy and grid.size > BIG):
            return []
        tabs = detect_tables(grid, off)
        # chỉ một bảng, bắt đầu ở A1 và cùng hình dạng với cách đọc thường -> không có gì mới, giữ luồng cũ
        if len(tabs) == 1 and tabs[0]["ref"].startswith("A1:") and tabs[0]["df"].shape == res["plain"].shape:
            return []
        return tabs
    except Exception:
        traceback.print_exc()
        return []


def _table_meta(t, n_tables, res):
    where = "Bảng Excel (Table)" if t["source"] == "excel_table" else "Bảng nhận diện từ bố cục"
    acts = [f"{where} «{t['name']}» tại {t['ref']}: {t['rows']:,} dòng × {t['cols']} cột".replace(",", ".")] + t["notes"]
    if n_tables > 1:
        acts.append(f"Ưu tiên #{t['rank']}/{n_tables} trong sheet" + (" (bảng chính)" if t["rank"] == 1 else ""))
    issues = list(res["issues"])[:3] if (res["verdict"] != "tidy" and t["rank"] == 1) else []   # nêu vấn đề của sheet một lần, ở bảng chính
    d = t["df"]
    pv = d.head(6)
    return {"verdict": "detected", "score": t["score"], "issues": issues, "actions": acts,
            "mode": "fixed", "ack": False, "range": t["ref"], "rank": t["rank"], "sheet_tables": n_tables,
            "preview": {"cols": [str(c) for c in pv.columns], "rows": json.loads(pv.to_json(orient="values", date_format="iso"))}}


def load_files(fs):
    """Đọc tệp + chạy cổng kiểm tra cấu trúc. Trả về (dfs đang dùng, báo cáo cấu trúc, bản gốc/bản chuẩn hóa).
    Sheet không chuẩn nhưng có bảng bên trong -> tách bảng ra dùng (xem tables.py); ngược lại giữ luồng cũ."""
    dfs, reports, alts = {}, {}, {}
    for f in fs:
        b = f.read()
        name = f.filename
        items = []
        if name.lower().endswith((".xlsx", ".xls", ".xlsm")):
            xl = pd.ExcelFile(io.BytesIO(b))
            wb = _open_wb(b) if name.lower().endswith((".xlsx", ".xlsm")) else None
            for sh in xl.sheet_names:
                key = f"{name}::{sh}" if len(xl.sheet_names) > 1 else name
                merged = len(wb[sh].merged_cells.ranges) if wb is not None else 0
                items.append((key, xl.parse(sh, header=None), xl.parse(sh), merged, ("xlsx" if wb is not None else "xls", wb, sh, None)))
        else:
            kw = dict(sep=None, engine="python", encoding="utf-8-sig", encoding_errors="replace")
            try:
                plain = pd.read_csv(io.BytesIO(b), **kw)
            except Exception:
                # không đọc được thành bảng (văn bản thuần, số cột lộn xộn...) -> vẫn nhận để cổng cấu trúc đánh giá
                lines = [l for l in b.decode("utf-8-sig", "replace").splitlines() if l.strip()]
                plain = pd.DataFrame({"noi_dung": lines})
            try:
                raw = pd.read_csv(io.BytesIO(b), header=None, **kw)
            except Exception:
                raw = None
            items.append((name, raw, plain, 0, ("csv", None, None, b)))
        for key, raw, plain, merged, src in items:
            res = analyze(raw, plain, merged)
            # báo cáo nhiều khối chồng nhau (tiêu đề lặp lại) đã có bộ chuẩn hóa riêng -> giữ nguyên luồng đó
            report_style = res["verdict"] == "fixed" and raw is not None and len(find_header_rows(raw)) >= 2
            tabs = [] if report_style else _found_tables(src, raw, res)
            if tabs:
                for t in tabs:
                    k = key if len(tabs) == 1 else f"{key}::{t['name']}"
                    dfs[k] = t["df"]
                    reports[k] = _table_meta(t, len(tabs), res)
                    alts[k] = {"plain": res["plain"], "fixed": t["df"]}
                continue
            dfs[key] = res["fixed"] if res["fixed"] is not None else res["plain"]
            reports[key] = meta(res)
            alts[key] = {"plain": res["plain"], "fixed": res["fixed"]}
    return dfs, reports, alts


def make_sample():
    rng = np.random.default_rng(7)
    n = 14500
    day = pd.Timestamp("2026-01-01") + pd.to_timedelta(rng.integers(0, 273, n), unit="D")
    reg = rng.choice(["HCM", "HN", "ĐN", "Khác"], n, p=[.42, .28, .14, .16])
    prod = rng.choice(["SP-A", "SP-B", "SP-C", "SP-D"], n, p=[.35, .3, .2, .15])
    base = pd.Series(prod).map({"SP-A": 1.0e6, "SP-B": .8e6, "SP-C": .6e6, "SP-D": .4e6}).to_numpy()
    q3 = np.asarray(day >= pd.Timestamp("2026-07-01"))
    hcm = reg == "HCM"
    price = base * rng.normal(1, .06, n) * np.where(q3 & hcm, .9, 1)
    keep = ~(q3 & hcm & (rng.random(n) < .14))
    d = pd.DataFrame({"ngay": day.strftime("%Y-%m-%d"), "ma_don": [f"O-{88000 + i}" for i in range(n)], "khu_vuc": reg,
                      "san_pham": prod, "customer_id": rng.integers(1, 2175, n), "so_luong": rng.integers(1, 8, n),
                      "don_gia": price.round(-3), "kenh": rng.choice(["Online", "Cửa hàng", "Đại lý"], n)})[keep].head(12480).reset_index(drop=True)
    d["doanh_thu"] = d["so_luong"] * d["don_gia"]
    d.loc[rng.choice(len(d), 1872, replace=False), "doanh_thu"] = np.nan
    d.loc[rng.choice(len(d), 42, replace=False), "ngay"] = "15/08/2026"
    d.loc[rng.choice(len(d), 9, replace=False), "doanh_thu"] = -500000
    d = pd.concat([d, d.sample(17, random_state=1)], ignore_index=True)
    cust = pd.DataFrame({"customer_id": range(1, 2141), "phan_khuc": rng.choice(["Lẻ", "Doanh nghiệp", "Đại lý"], 2140),
                         "thanh_pho": rng.choice(["HCM", "HN", "ĐN", "Khác"], 2140)})
    prods = pd.DataFrame({"san_pham": ["SP-A", "SP-B", "SP-C", "SP-D"], "nhom": ["Chủ lực", "Chủ lực", "Phổ thông", "Phụ kiện"]})
    return {"sales_q3.csv": d, "customers.csv": cust, "products.csv": prods}